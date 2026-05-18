# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from outbound_validation_spec import (
    CUSTOMER_ROWS,
    DAY_ROWS,
    FINAL_DECISION,
    MICRO_DIAGNOSIS_ROWS,
    OUTREACH_TEMPLATE,
    PHONE_QUESTIONS,
    PRODUCT_ROWS,
    REPLY_RULE_ROWS,
    REVIEW_ROWS,
    SCORING_ROWS,
    STARTER_CHECKLIST,
    SUBTITLE,
    SUCCESS_ROWS,
    THESIS_BODY,
    THESIS_TITLE,
    TITLE,
)


OUT = Path(__file__).with_name("出海销售表达服务_7天付费验证执行稿.docx")

BLUE = "163D5C"
BLUE_2 = "255C7E"
PALE = "EAF2F8"
GRAY = "F4F6F8"
BORDER = "D7E1EA"
TEXT = "1F2933"
MUTED = "52616B"
WHITE = "FFFFFF"
RED = "8A1F11"


def east_asia(run, font="Microsoft YaHei"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=115, start=115, bottom=115, end=115):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        elem = tbl_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tbl_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def cell_text(cell, text, bold=False, color=TEXT, size=9.0, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    east_asia(run)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    for row in table.rows:
        keep_row(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                for r in p.runs:
                    r.font.size = Pt(9.0)
                    r.font.color.rgb = RGBColor.from_string(TEXT)
                    east_asia(r)


def table(doc, headers, rows, widths=None, header_fill=BLUE):
    tbl = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        c = tbl.rows[0].cells[i]
        cell_text(c, header, bold=True, color=WHITE, size=9.1)
        shade(c, header_fill)
        if widths:
            c.width = widths[i]
    for row in rows:
        cells = tbl.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, bold=(i == 0), size=8.9)
            if i == 0:
                shade(cells[i], GRAY)
            if widths:
                cells[i].width = widths[i]
    style_table(tbl)
    for cell in tbl.rows[0].cells:
        shade(cell, header_fill)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(WHITE)
                east_asia(r)
    doc.add_paragraph()
    return tbl


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)


def para(doc, text, size=10.0, color=TEXT, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    east_asia(run)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(item)
        r.font.size = Pt(9.4)
        r.font.color.rgb = RGBColor.from_string(TEXT)
        east_asia(r)


def callout(doc, title, body, fill=PALE, color=BLUE):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(tbl, color=fill)
    c = tbl.cell(0, 0)
    shade(c, fill)
    margins(c, 150, 170, 150, 170)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(color)
    east_asia(r)
    p2 = c.add_paragraph()
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(body)
    r2.font.size = Pt(9.6)
    r2.font.color.rgb = RGBColor.from_string(TEXT)
    east_asia(r2)
    doc.add_paragraph()


def page(doc):
    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        p = doc.paragraphs[-1]._element
        p.getparent().remove(p)
    doc.add_page_break()


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.55)
    sec.bottom_margin = Cm(1.45)
    sec.left_margin = Cm(1.65)
    sec.right_margin = Cm(1.65)
    sec.footer_distance = Cm(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)

    for name in ["Heading 1", "Heading 2", "Title", "Subtitle"]:
        s = styles[name]
        s.font.name = "Microsoft YaHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.color.rgb = RGBColor.from_string(BLUE)
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.size = Pt(11)
    styles["Heading 2"].font.bold = True
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Subtitle"].font.size = Pt(10)
    styles["Subtitle"].font.color.rgb = RGBColor.from_string(MUTED)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("7 天付费验证执行单")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    east_asia(r)


def build():
    doc = Document()
    setup(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run(TITLE)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(SUBTITLE)

    callout(doc, THESIS_TITLE, THESIS_BODY)

    h2(doc, "今天先做这 5 件事")
    bullets(doc, STARTER_CHECKLIST)

    table(
        doc,
        ["7 天唯一成功标准", "不算成功"],
        SUCCESS_ROWS,
        widths=[Cm(8.1), Cm(8.1)],
    )

    h1(doc, "1. 只卖这一个东西")
    table(
        doc,
        ["产品", "价格", "交付物", "边界"],
        PRODUCT_ROWS,
        widths=[Cm(3.8), Cm(1.9), Cm(7.2), Cm(3.3)],
    )
    page(doc)
    h1(doc, "2. 只找这类客户")
    table(
        doc,
        ["优先级", "客户类型", "可见信号", "判断"],
        CUSTOMER_ROWS,
        widths=[Cm(1.5), Cm(4.1), Cm(5.2), Cm(5.4)],
    )

    h2(doc, "客户筛选打分")
    table(
        doc,
        ["维度", "2 分", "1 分", "0 分"],
        SCORING_ROWS,
        widths=[Cm(2.7), Cm(5.0), Cm(4.4), Cm(4.1)],
    )
    para(doc, "规则：总分 7 分以上才做微诊断；低于 7 分不投入时间。", bold=True, color=BLUE)

    page(doc)
    h1(doc, "3. 7 天执行计划")
    table(
        doc,
        ["天", "当天目标", "必须完成", "通过/失败标准", "禁止"],
        DAY_ROWS,
        widths=[Cm(0.9), Cm(2.5), Cm(5.3), Cm(4.5), Cm(3.0)],
    )

    page(doc)
    h1(doc, "4. 触达和电话话术")
    h2(doc, "定制触达模板")
    callout(
        doc,
        "私信/微信/邮件",
        OUTREACH_TEMPLATE,
        fill=GRAY,
    )
    h2(doc, "电话必须问")
    bullets(doc, PHONE_QUESTIONS)

    h2(doc, "客户回复后这样处理")
    table(
        doc,
        ["客户反应", "处理方式"],
        REPLY_RULE_ROWS,
        widths=[Cm(4.5), Cm(11.7)],
    )

    h2(doc, "微诊断只写这 4 行")
    table(
        doc,
        ["行", "填写内容"],
        MICRO_DIAGNOSIS_ROWS,
        widths=[Cm(3.4), Cm(12.8)],
    )

    page(doc)
    h1(doc, "5. 日终复盘表")
    table(
        doc,
        ["指标", "目标", "实际", "判断"],
        REVIEW_ROWS,
        widths=[Cm(4.1), Cm(3.2), Cm(2.5), Cm(6.4)],
    )

    callout(
        doc,
        "最终判断",
        FINAL_DECISION,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
